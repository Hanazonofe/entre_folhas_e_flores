import io

import pytest
from openpyxl import Workbook, load_workbook
from pdv import models as m
from pdv import product_spreadsheet as sheet
from sqlalchemy import event
from sqlalchemy.exc import IntegrityError


def book(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.append(headers or sheet.HEADERS)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(client, data, **params):
    return client.post(
        "/api/products/import",
        params=params,
        content=data,
        headers={"Content-Type": sheet.MIME},
    )


# @principle:P-005
def test_operator_export_preview_apply_and_conflict(client, clean):
    login = client.post(
        "/api/auth/login", json={"login": "operator", "password": "test-password-123"}
    )
    client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
    p = client.post(
        "/api/products",
        json={
            "code": "001",
            "name": "=Literal",
            "price_cents": 1250,
            "barcode": "000123",
            "active": False,
        },
    ).json()
    response = client.get("/api/products/export")
    assert (
        response.status_code == 200 and response.headers["content-type"] == sheet.MIME
    )
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws["A2"].value == "001" and ws["D2"].value == "000123"
    assert ws["B2"].value == "=Literal" and ws["B2"].data_type == "s"
    assert client.get("/api/products").json()["count"] == 1
    assert client.get("/api/products?active_only=true").json()["count"] == 0
    ws["B2"] = "Alterado"
    ws["C2"] = "R$ 23,45"
    ws["E2"] = "1,234"
    ws.append(["002", "Novo", "10.00", "", "0", "sim"])
    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    preview = upload(client, data).json()
    assert (preview["criar"], preview["atualizar"], preview["invalidos"]) == (1, 1, 0)
    assert client.get("/api/products").json()["items"][0]["name"] == "=Literal"
    applied = upload(client, data, apply=True, preview=preview["preview"])
    assert applied.status_code == 200, applied.text
    assert applied.json()["inseridos"] == 1 and applied.json()["atualizados"] == 1
    result = {p["code"]: p for p in client.get("/api/products").json()["items"]}
    assert result["001"]["price_cents"] == 2345 and result["001"]["stock"] == 1.234
    assert result["002"]["barcode"] is None and result["001"]["id"] == p["id"]
    assert (
        upload(client, data, apply=True, preview=preview["preview"]).status_code == 409
    )
    edit = {
        k: result["002"][k]
        for k in (
            "code",
            "barcode",
            "name",
            "price_cents",
            "stock",
            "active",
            "version",
        )
    }
    edit["name"] = "Individual"
    assert (
        client.put("/api/products/" + result["002"]["id"], json=edit).status_code == 200
    )
    assert client.get("/api/users").status_code == 403


def test_invalid_skip_and_optional_preservation(client):
    client.post(
        "/api/products",
        json={
            "code": "A",
            "name": "A",
            "price_cents": 10,
            "stock": 2,
            "barcode": "001",
            "active": False,
        },
    ).json()
    data = book(
        [
            ["A", "Edited", "20"],
            ["B", "Valid", "30"],
            ["B", "Duplicate", "40"],
            ["C", "", "10"],
        ],
        ["CODIGO", "PRODUTO", "VALOR"],
    )
    preview = upload(client, data).json()
    assert preview["invalidos"] == 3 and preview["validos"] == 1
    assert (
        upload(client, data, apply=True, preview=preview["preview"]).status_code == 422
    )
    preview = upload(client, data, skip_invalid=True).json()
    assert (
        upload(
            client, data, apply=True, skip_invalid=True, preview=preview["preview"]
        ).status_code
        == 200
    )
    result = client.get("/api/products").json()["items"][0]
    assert (
        result["stock"] == 2
        and result["barcode"] == "001"
        and result["active"] is False
    )


@pytest.mark.parametrize(
    "row",
    [
        ["A", "Formula", "=1+1", "", "0", "sim"],
        ["A", "Precision", "1.001", "", "0", "sim"],
        ["A", "Stock", "1", "", "0.0001", "sim"],
        ["A", "Status", "1", "", "0", "maybe"],
        ["A", "Identifier", "1", 123, "0", "sim"],
    ],
)
# @principle:P-005
def test_validation(client, row):
    result = upload(client, book([row]))
    assert result.status_code == 200 and result.json()["invalidos"] == 1


def test_barcode_conflict_and_atomicity(client):
    client.post(
        "/api/products",
        json={"code": "A", "name": "A", "price_cents": 10, "barcode": "001"},
    )
    data = book([["B", "New", "20", "001", "0", "sim"]])
    assert upload(client, data).json()["invalidos"] == 1
    data = book(
        [["B", "New", "20", "", "0", "sim"], ["A", "Edited", "30", "", "0", "sim"]]
    )
    preview = upload(client, data).json()

    def fail(*args, **kwargs):
        raise IntegrityError("injected", {}, Exception())

    # UPDATE is emitted after pending changes have reached PostgreSQL; failure rolls everything back.
    event.listen(m.Product, "after_update", fail)
    try:
        assert (
            upload(client, data, apply=True, preview=preview["preview"]).status_code
            == 409
        )
    finally:
        event.remove(m.Product, "after_update", fail)
    products = client.get("/api/products").json()["items"]
    assert len(products) == 1 and products[0]["name"] == "A"


def test_stale_preview_and_security(client):
    data = book([["A", "New", "20", "", "0", "sim"]])
    preview = upload(client, data).json()
    client.post("/api/products", json={"code": "B", "name": "B", "price_cents": 10})
    assert (
        upload(client, data, apply=True, preview=preview["preview"]).status_code == 409
    )
    assert (
        client.post(
            "/api/products/import",
            content=data,
            headers={"Content-Type": sheet.MIME, "X-CSRF-Token": "wrong"},
        ).status_code
        == 403
    )
    assert upload(client, b"broken").status_code == 422
    assert upload(client, b"x" * 5_242_881).status_code == 413
    client.cookies.clear()
    assert client.get("/api/products/export").status_code == 401
    assert upload(client, data).status_code == 401


# @principle:P-005
def test_export_exact_money_and_empty(client):
    data = client.get("/api/products/export").content
    assert load_workbook(io.BytesIO(data)).active.max_row == 1
    client.post(
        "/api/products",
        json={"code": "001", "name": "Max", "price_cents": 9007199254740991},
    )
    data = client.get("/api/products/export").content
    parsed, errors = sheet.read_xlsx(data)
    assert not errors and parsed[0]["values"]["price_cents"] == 9007199254740991
